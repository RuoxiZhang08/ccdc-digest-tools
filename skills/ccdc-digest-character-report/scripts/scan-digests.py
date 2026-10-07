#!/usr/bin/env python3
"""Scan CCDC digest workbooks for characters rejected by the CCDC ETL rule.

This is the Python counterpart to scan-digests.mjs. It reads source workbooks
without modifying them and writes a separate Excel report.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import warnings
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


warnings.filterwarnings(
    "ignore",
    message="Data Validation extension is not supported and will be removed",
)


REPORT_FOLDER = "ccdc-character-reports"
MAX_EXCEL_TEXT = 32_767

# ECMAScript String.prototype.trim() whitespace and line terminators.
JS_TRIM_CHARACTERS = (
    "\u0009\u000A\u000B\u000C\u000D\u0020\u00A0\u1680"
    "\u2000\u2001\u2002\u2003\u2004\u2005\u2006\u2007\u2008\u2009\u200A"
    "\u2028\u2029\u202F\u205F\u3000\uFEFF"
)
ILLEGAL_XML_CHARACTER = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")

RESOURCE_FIELDS = [
    (10, 1, "Initial Submission Date", "initial_submission_date"),
    (10, 3, "Resource ID", "id"),
    (10, 4, "Resource Name", "resource_name"),
    (10, 5, "Resource URI", "resource_uri"),
    (15, 1, "Resource Description", "description"),
    (15, 2, "Resource Type", "resource_type"),
    (15, 3, "Data Content Type", "data_content_type"),
    (15, 4, "Resource POC", "poc"),
    (15, 5, "Resource POC Email", "poc_email"),
    (15, 6, "Analytics", "analytics"),
    (15, 7, "Visualization", "visualization"),
    (15, 8, "Pediatric Specific", "pediatric_specific"),
    (15, 9, "API", "api"),
    (18, 1, "Data Update Date", "data_update_date"),
    (18, 5, "Resource Contact URL", "resource_contact_url"),
]

DATASET_FIELDS = [
    (3, "Dataset Abbreviated Name", "dataset_name"),
    (4, "Dataset Full Name", "dataset_full_name"),
    (5, "Dataset Description", "description"),
    (6, "Primary Dataset Scope", "dataset_scope"),
    (7, "Dataset POC", "poc"),
    (8, "POC Email", "poc_email"),
    (9, "Reference", "published_in"),
]

DIGEST_FIELDS = [
    (2, "Parent Digest ID", "parent_digest_level"),
    (3, "Digest ID", "digest_level"),
    (4, "Core Element", "core_element"),
    (5, "Data Element", "data_element"),
    (6, "Data Element Value", "element_value"),
    (7, "Statistic Type", "statistic_type"),
    (8, "Statistic Value", "statistic_value"),
]


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Create a CCDC digest special-character report."
    )
    parser.add_argument("--input", required=True, help="Folder containing digest .xlsx files")
    parser.add_argument("--output", help="Optional output .xlsx path")
    return parser.parse_args(argv)


def timestamp(value: datetime | None = None) -> str:
    current = value or datetime.now(timezone.utc)
    rendered = current.astimezone(timezone.utc).isoformat(timespec="milliseconds")
    return rendered.replace("+00:00", "Z").replace(":", "-").replace(".", "-")


def readable_timestamp(value: datetime) -> str:
    return value.astimezone(timezone.utc).strftime("%m/%d/%Y, %H:%M:%S UTC")


def js_string(value: Any) -> str:
    if value is None:
        return ""
    if value is True:
        return "true"
    if value is False:
        return "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return str(value)
    return str(value)


def normalized_value(value: Any) -> str:
    return js_string(value).strip(JS_TRIM_CHARACTERS)


def is_allowed_character(character: str) -> bool:
    return character in {"\r", "\n"} or "\x20" <= character <= "\x7e"


def invalid_occurrences(value: Any) -> tuple[str, list[dict[str, Any]]]:
    normalized = normalized_value(value)
    if not normalized or all(is_allowed_character(character) for character in normalized):
        return normalized, []
    occurrences: list[dict[str, Any]] = []
    for position, character in enumerate(normalized, start=1):
        if is_allowed_character(character):
            continue
        occurrences.append(
            {
                "character": character,
                "codePoint": f"U+{ord(character):04X}",
                "position": position,
            }
        )
    return normalized, occurrences


def display_character(character: str) -> str:
    names = {
        "\t": "TAB",
        "\u00A0": "NO-BREAK SPACE",
        "\u200B": "ZERO WIDTH SPACE",
    }
    return names.get(character, character)


def clean_value(normalized: str) -> str:
    return "".join(character for character in normalized if is_allowed_character(character))


def excel_safe_string(value: str) -> str:
    return ILLEGAL_XML_CHARACTER.sub(
        lambda match: f"\\u{ord(match.group(0)):04X}", value
    )


def safe_text(value: Any) -> str:
    result = js_string(value)
    if len(result) > MAX_EXCEL_TEXT:
        result = f"{result[: MAX_EXCEL_TEXT - 14]}… [truncated]"
    if result.startswith(("=", "+", "-", "@")):
        result = f"'{result}"
    return excel_safe_string(result)


def value_at(data: list[tuple[Any, ...]], one_based_row: int, one_based_column: int) -> Any:
    row_index = one_based_row - 1
    column_index = one_based_column - 1
    if row_index < 0 or row_index >= len(data):
        return None
    row = data[row_index]
    if column_index < 0 or column_index >= len(row):
        return None
    return row[column_index]


def row_has_value(row: tuple[Any, ...] | None) -> bool:
    return bool(row) and any(value not in (None, "") for value in row)


def worksheet_data(sheet: Any) -> list[tuple[Any, ...]]:
    return list(sheet.iter_rows(values_only=True))


def make_finding(
    *,
    filename: str,
    sheet_name: str,
    sheet_index: int,
    row: int,
    column: int,
    field: str,
    etl_field: str,
    value: Any,
    dataset_name: Any = "",
) -> dict[str, Any] | None:
    normalized, occurrences = invalid_occurrences(value)
    if not occurrences:
        return None

    distinct: list[str] = []
    seen: set[tuple[str, str]] = set()
    for occurrence in occurrences:
        key = (occurrence["character"], occurrence["codePoint"])
        if key in seen:
            continue
        seen.add(key)
        distinct.append(
            f"{display_character(occurrence['character'])} ({occurrence['codePoint']})"
        )

    column_name = get_column_letter(column)
    return {
        "filename": filename,
        "sheetName": sheet_name,
        "sheetIndex": sheet_index,
        "datasetName": dataset_name,
        "row": row,
        "column": column_name,
        "cell": f"{column_name}{row}",
        "field": field,
        "etlField": etl_field,
        "originalValue": safe_text(value),
        "specialCharacters": safe_text(", ".join(distinct)),
        "unicodeCodePoints": ", ".join(item["codePoint"] for item in occurrences),
        "characterPositions": ", ".join(str(item["position"]) for item in occurrences),
        "cleanedValue": safe_text(clean_value(normalized)),
        "occurrenceCount": len(occurrences),
    }


def select_digest_sheets(sheets: list[Any], dataset_names: list[Any]) -> list[Any | None]:
    digest_sheets = sheets[2:]
    selected: list[Any | None] = [None] * len(dataset_names)
    claimed: set[int] = set()

    for index, dataset_name in enumerate(dataset_names):
        if not dataset_name:
            continue
        named = next(
            (
                sheet
                for sheet in digest_sheets
                if sheet.title == dataset_name and id(sheet) not in claimed
            ),
            None,
        )
        if named is not None:
            selected[index] = named
            claimed.add(id(named))

    for index, dataset_name in enumerate(dataset_names):
        if not dataset_name or selected[index] is not None:
            continue
        available = digest_sheets[index] if index < len(digest_sheets) else None
        if available is not None and id(available) in claimed:
            unclaimed = [sheet for sheet in digest_sheets if id(sheet) not in claimed]
            if len(unclaimed) != 1:
                raise ValueError(f"Digest sheet for {dataset_name} is ambiguous")
            available = unclaimed[0]
        if available is None:
            raise ValueError(f"Digest sheet for {dataset_name} is missing")
        selected[index] = available
        claimed.add(id(available))

    return selected


def scan_workbook(workbook: Any, filename: str) -> dict[str, Any]:
    sheets = list(workbook.worksheets)
    findings: list[dict[str, Any]] = []
    errors: list[str] = []

    def add(**details: Any) -> None:
        finding = make_finding(filename=filename, **details)
        if finding is not None:
            findings.append(finding)

    if not sheets:
        errors.append("Missing Data Resource sheet")
    else:
        resource = sheets[0]
        resource_data = worksheet_data(resource)
        for row, column, field, etl_field in RESOURCE_FIELDS:
            add(
                sheet_name=resource.title,
                sheet_index=0,
                row=row,
                column=column,
                field=field,
                etl_field=etl_field,
                value=value_at(resource_data, row, column),
            )

    if len(sheets) < 2:
        errors.append("Missing Dataset Information sheet")
        return {"findings": findings, "errors": errors, "sheetCount": len(sheets)}

    dataset = sheets[1]
    dataset_data = worksheet_data(dataset)
    for row_number in range(2, len(dataset_data) + 1):
        dataset_name = value_at(dataset_data, row_number, 3) or f"Dataset_{row_number - 1}"
        for column, field, etl_field in DATASET_FIELDS:
            add(
                sheet_name=dataset.title,
                sheet_index=1,
                dataset_name=dataset_name,
                row=row_number,
                column=column,
                field=field,
                etl_field=etl_field,
                value=value_at(dataset_data, row_number, column),
            )

    dataset_names: list[Any] = []
    for index, row in enumerate(dataset_data[1:], start=1):
        if not row_has_value(row):
            dataset_names.append(None)
        else:
            dataset_names.append((row[2] if len(row) > 2 else None) or f"Dataset_{index}")

    try:
        selected = select_digest_sheets(sheets, dataset_names)
    except Exception as error:
        errors.append(str(error))
        return {"findings": findings, "errors": errors, "sheetCount": len(sheets)}

    for index, dataset_name in enumerate(dataset_names):
        if not dataset_name:
            continue
        sheet = selected[index]
        if sheet is None:
            errors.append(f"Digest sheet for {dataset_name} is missing")
            continue
        data = worksheet_data(sheet)
        for row_number in range(2, len(data) + 1):
            for column, field, etl_field in DIGEST_FIELDS:
                add(
                    sheet_name=sheet.title,
                    sheet_index=sheets.index(sheet),
                    dataset_name=dataset_name,
                    row=row_number,
                    column=column,
                    field=field,
                    etl_field=etl_field,
                    value=value_at(data, row_number, column),
                )

    return {"findings": findings, "errors": errors, "sheetCount": len(sheets)}


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
LABEL_FILL = PatternFill("solid", fgColor="D9EAF7")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10, color="222222")
LABEL_FONT = Font(name="Arial", size=10, bold=True, color="222222")
TITLE_FONT = Font(name="Arial", size=14, bold=True, color="1F4E78")
WHITE_SIDE = Side(style="thin", color="FFFFFF")
HEADER_BORDER = Border(left=WHITE_SIDE, right=WHITE_SIDE)


def append_rows(sheet: Any, rows: Iterable[Iterable[Any]]) -> None:
    for row in rows:
        sheet.append(list(row))


def write_table_sheet(
    workbook: Workbook,
    name: str,
    headers: list[str],
    rows: list[list[Any]],
    widths: list[float],
    table_name: str,
) -> Any:
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    append_rows(sheet, [headers, *rows])

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = HEADER_BORDER

    if rows:
        for row in sheet.iter_rows(min_row=2, max_row=len(rows) + 1, max_col=len(headers)):
            for cell in row:
                cell.font = BODY_FONT
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        table = Table(
            displayName=table_name,
            ref=f"A1:{get_column_letter(len(headers))}{len(rows) + 1}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

    sheet.freeze_panes = "A2"
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    return sheet


def write_report(
    *,
    input_path: Path,
    output_path: Path,
    run_at: datetime,
    files: list[dict[str, Any]],
    findings: list[dict[str, Any]],
    errors: list[dict[str, str]],
) -> dict[str, int]:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.sheet_view.showGridLines = False

    affected_files = len({item["filename"] for item in findings})
    occurrences = sum(item["occurrenceCount"] for item in findings)
    scanned = sum(item["status"] != "ERROR" for item in files)
    summary_rows = [
        ["CCDC Digest Character Report", ""],
        ["Input folder", safe_text(input_path)],
        ["Scan time", readable_timestamp(run_at)],
        ["Workbooks found", len(files)],
        ["Workbooks scanned", scanned],
        ["Workbooks with findings", affected_files],
        ["Workbooks with errors", sum(item["status"] == "ERROR" for item in files)],
        ["Affected cells", len(findings)],
        ["Invalid-character occurrences", occurrences],
        ["Character rule", "After trim, allow CR, LF, and printable ASCII U+0020–U+007E."],
        ["Clean-value rule", "Delete rejected code points from the trimmed value. Source workbooks are unchanged."],
    ]
    append_rows(summary, summary_rows)
    for cell in summary[1]:
        cell.font = TITLE_FONT
    for row in range(2, 12):
        summary.cell(row, 1).font = LABEL_FONT
        summary.cell(row, 1).fill = LABEL_FILL
        summary.cell(row, 1).alignment = Alignment(vertical="top")
        summary.cell(row, 2).font = BODY_FONT
        summary.cell(row, 2).alignment = Alignment(vertical="top", wrap_text=True)
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 74

    finding_rows = [
        [
            safe_text(item["filename"]),
            safe_text(item["sheetName"]),
            item["sheetIndex"],
            safe_text(item["datasetName"]),
            item["row"],
            item["column"],
            item["cell"],
            item["field"],
            item["etlField"],
            item["originalValue"],
            item["specialCharacters"],
            item["unicodeCodePoints"],
            item["characterPositions"],
            item["cleanedValue"],
            item["occurrenceCount"],
        ]
        for item in findings
    ]
    write_table_sheet(
        workbook,
        "Special Characters",
        [
            "Filename",
            "Sheet",
            "Sheet Index",
            "Dataset",
            "Row",
            "Column",
            "Cell",
            "Field",
            "ETL Field",
            "Original Value",
            "Special Characters",
            "Unicode Code Points",
            "Character Positions",
            "Value Without Special Characters",
            "Occurrence Count",
        ],
        finding_rows,
        [28, 24, 11, 20, 9, 9, 11, 24, 22, 42, 28, 25, 20, 42, 12],
        "SpecialCharactersTable",
    )

    file_rows = [
        [
            safe_text(item["filename"]),
            item["status"],
            item["sheetCount"] if item["sheetCount"] is not None else "",
            item["affectedCells"],
            item["occurrenceCount"],
            safe_text(item["errorMessage"]),
        ]
        for item in files
    ]
    write_table_sheet(
        workbook,
        "Files",
        [
            "Filename",
            "Status",
            "Sheet Count",
            "Affected Cells",
            "Invalid-character Occurrences",
            "Error Message",
        ],
        file_rows,
        [30, 13, 13, 14, 25, 58],
        "FilesTable",
    )

    error_rows = [
        [safe_text(item["filename"]), item["stage"], safe_text(item["message"])]
        for item in errors
    ]
    write_table_sheet(
        workbook,
        "Errors",
        ["Filename", "Stage", "Message"],
        error_rows,
        [30, 20, 70],
        "ErrorsTable",
    )

    write_table_sheet(
        workbook,
        "Rules",
        ["Rule", "Definition", "Source"],
        [
            [
                "CCDC ETL character rule",
                "Convert the value to text, trim it, allow blank, CR, LF, and printable ASCII U+0020–U+007E; report every other code point.",
                "CCDC-ETL/common/utils.js containsSpecialCharacters",
            ],
            [
                "Value Without Special Characters",
                "Delete rejected code points from the trimmed value. This report-only suggestion is not an ETL transformation.",
                "CCDC Digest Character Report",
            ],
        ],
        [34, 90, 52],
        "RulesTable",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return {
        "affectedFiles": affected_files,
        "occurrences": occurrences,
        "scanned": scanned,
    }


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv if argv is not None else sys.argv[1:])
    input_path = Path(args.input).expanduser().resolve()
    if not input_path.exists():
        raise ValueError(f"Cannot access input folder {input_path}: path does not exist")
    if not input_path.is_dir():
        raise ValueError(f"Input path is not a folder: {input_path}")

    output_path = (
        Path(args.output).expanduser().resolve()
        if args.output
        else input_path
        / REPORT_FOLDER
        / f"special-character-report-{timestamp()}.xlsx"
    )
    if output_path.suffix.lower() != ".xlsx":
        raise ValueError("Output path must end in .xlsx")

    filenames = sorted(
        (
            item.name
            for item in input_path.iterdir()
            if item.is_file()
            and not item.name.startswith(".")
            and not item.name.startswith("~$")
            and item.suffix.lower() == ".xlsx"
        ),
        key=lambda name: (name.casefold(), name),
    )
    source_paths = {(input_path / filename).resolve() for filename in filenames}
    if output_path.resolve() in source_paths:
        raise ValueError("Output path must not overwrite an input workbook")

    files: list[dict[str, Any]] = []
    findings: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []

    for filename in filenames:
        file_path = input_path / filename
        workbook = None
        try:
            workbook = load_workbook(file_path, read_only=True, data_only=True)
            result = scan_workbook(workbook, filename)
            findings.extend(result["findings"])
            errors.extend(
                {
                    "filename": filename,
                    "stage": "Workbook mapping",
                    "message": message,
                }
                for message in result["errors"]
            )
            occurrence_count = sum(
                item["occurrenceCount"] for item in result["findings"]
            )
            files.append(
                {
                    "filename": filename,
                    "status": (
                        "ERROR"
                        if result["errors"]
                        else "FINDINGS"
                        if result["findings"]
                        else "CLEAN"
                    ),
                    "sheetCount": result["sheetCount"],
                    "affectedCells": len(result["findings"]),
                    "occurrenceCount": occurrence_count,
                    "errorMessage": "; ".join(result["errors"]),
                }
            )
        except Exception as error:
            message = str(error)
            errors.append(
                {"filename": filename, "stage": "Workbook read", "message": message}
            )
            files.append(
                {
                    "filename": filename,
                    "status": "ERROR",
                    "sheetCount": None,
                    "affectedCells": 0,
                    "occurrenceCount": 0,
                    "errorMessage": message,
                }
            )
        finally:
            if workbook is not None:
                workbook.close()

    findings.sort(
        key=lambda item: (
            item["filename"],
            item["sheetIndex"],
            item["row"],
            item["column"],
        )
    )

    run_at = datetime.now(timezone.utc)
    totals = write_report(
        input_path=input_path,
        output_path=output_path,
        run_at=run_at,
        files=files,
        findings=findings,
        errors=errors,
    )
    summary = {
        "ok": True,
        "inputPath": str(input_path),
        "outputPath": str(output_path),
        "workbooksFound": len(files),
        "workbooksScanned": totals["scanned"],
        "workbooksWithFindings": totals["affectedFiles"],
        "workbooksWithErrors": sum(item["status"] == "ERROR" for item in files),
        "affectedCells": len(findings),
        "invalidCharacterOccurrences": totals["occurrences"],
    }
    print(json.dumps(summary, ensure_ascii=False))
    return 1 if findings or errors else 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except Exception as error:
        print(json.dumps({"ok": False, "error": str(error)}, ensure_ascii=False), file=sys.stderr)
        raise SystemExit(2)
